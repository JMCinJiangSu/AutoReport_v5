import os
import pandas as pd
from docx import Document

# 配置参数：Word文件所在文件夹路径（请修改为你的实际路径）
WORD_FOLDER = r"mnt\d\Scripts\AutoReport_v5\test"
# 输出Excel文件路径
EXCEL_OUTPUT = r"mnt\d\Scripts\AutoReport_v5\test\WFY_output.xlsx"

# 初始化结果列表
result_data = []

# 遍历文件夹中所有Word文件
for filename in os.listdir(WORD_FOLDER):
    if filename.endswith(".docx"):  # 仅处理docx文件
        file_path = os.path.join(WORD_FOLDER, filename)
        try:
            # 打开Word文档
            doc = Document(file_path)
            
            # 提取姓名：第一行文本（示例格式：东阳市人民医院22-639B）
            if doc.paragraphs:
                name = doc.paragraphs[0].text.strip()
            else:
                print(f"文件{filename}无首行文本，跳过")
                continue
            
            # 提取表格数据（假设Word中只有1个目标表格）
            if not doc.tables:
                print(f"文件{filename}无表格，跳过")
                continue
            target_table = doc.tables[0]
            
            # 遍历表格行（跳过表头行，从数据行开始）
            # 假设表格列顺序：基因 | 变异 | 突变型 | 丰度/拷贝数
            for row_idx, row in enumerate(target_table.rows):
                if row_idx == 0:  # 跳过表头（如表格有表头行）
                    continue
                # 提取每行单元格内容
                cells = row.cells
                if len(cells) >= 4:  # 确保有4列数据
                    gene = cells[0].text.strip()
                    variation = cells[1].text.strip()
                    mutation_type = cells[2].text.strip()
                    abundance = cells[3].text.strip()
                    # 追加到结果列表（姓名复用）
                    result_data.append({
                        "姓名": name,
                        "基因": gene,
                        "变异": variation,
                        "突变型": mutation_type,
                        "丰度/拷贝数": abundance
                    })
            print(f"文件{filename}处理完成")
        except Exception as e:
            print(f"处理文件{filename}出错：{str(e)}")

# 将结果转为DataFrame并生成Excel
if result_data:
    df = pd.DataFrame(result_data)
    # 保存Excel（openpyxl为xlsx格式引擎）
    df.to_excel(EXCEL_OUTPUT, index=False, engine="openpyxl")
    print(f"所有文件处理完成！结果已保存至：{EXCEL_OUTPUT}")
else:
    print("未提取到任何有效数据！")